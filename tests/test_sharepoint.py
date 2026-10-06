"""RepositorioSharePoint con un SharePoint simulado en memoria, y AlmacenGraph contra un
Microsoft Graph simulado (sin red)."""

import io
from datetime import datetime

import pytest
from acta_app.storage.excel_formato import leer_hipervinculo
from openpyxl import load_workbook

from acta_app.models import Articulo
from acta_app.storage import ActaDuplicadaError, AlmacenamientoError
from acta_app.storage.sharepoint import (
    AlmacenGraph,
    Archivo,
    ConflictoDeVersion,
    RepositorioSharePoint,
)

CARPETA = "16. Analisis de Datos/Actas"


class SharePointFalso:
    """Almacén en memoria con eTags, como una biblioteca de SharePoint."""

    def __init__(self):
        self.archivos: dict[str, tuple[bytes, int]] = {}
        self.carpetas = {CARPETA}
        self.antes_de_escribir = None  # para simular a otro usuario guardando

    def _url(self, ruta):
        return f"https://sp.example/{ruta}"

    def leer(self, ruta):
        if ruta not in self.archivos:
            return None
        datos, version = self.archivos[ruta]
        return Archivo(datos=datos, version=f'"v{version}"', enlace=self._url(ruta))

    def escribir(self, ruta, datos, version_esperada=None, solo_si_no_existe=False):
        if self.antes_de_escribir and ruta.endswith(".xlsx"):
            accion, self.antes_de_escribir = self.antes_de_escribir, None
            accion()
        actual = self.archivos.get(ruta)
        if solo_si_no_existe and actual is not None:
            raise ConflictoDeVersion(ruta)
        if version_esperada and (actual is None or f'"v{actual[1]}"' != version_esperada):
            raise ConflictoDeVersion(ruta)
        self.archivos[ruta] = (datos, (actual[1] + 1) if actual else 1)
        return self._url(ruta)

    def eliminar(self, ruta):
        self.archivos.pop(ruta, None)

    def listar(self, ruta):
        hijos = {}
        for r in [*self.archivos, *self.carpetas]:
            if r.startswith(ruta + "/"):
                nombre = r[len(ruta) + 1:].split("/")[0]
                hijos[nombre] = hijos.get(nombre) or f"{ruta}/{nombre}" in self.carpetas or (
                    "/" in r[len(ruta) + 1:])
        if not hijos and ruta not in self.carpetas:
            return None
        return sorted(hijos.items())

    def enlace(self, ruta):
        existe = ruta in self.archivos or ruta in self.carpetas or any(
            r.startswith(ruta + "/") for r in self.archivos
        )
        return self._url(ruta) if existe else None


@pytest.fixture
def sp():
    return SharePointFalso()


@pytest.fixture
def repo(sp):
    return RepositorioSharePoint(sp, carpeta=CARPETA, segundos_cache=0)


def _fila(sp, n=1):
    ws = load_workbook(io.BytesIO(sp.archivos[f"{CARPETA}/Actas.xlsx"][0]))["Actas"]
    return {ws.cell(2, c).value: ws.cell(2 + n, c) for c in range(1, ws.max_column + 1)}


def test_primera_acta_crea_actas_xlsx_y_sube_pdf_y_firmas(repo, sp, acta_completa):
    resultado = repo.guardar(acta_completa, b"%PDF", "Acta_2026-00051.pdf")

    assert resultado.total_actas == 1
    assert resultado.ubicacion_pdf == f"https://sp.example/{CARPETA}/PDF/Acta_2026-00051.pdf"
    assert sp.archivos[f"{CARPETA}/PDF/Acta_2026-00051.pdf"][0] == b"%PDF"
    assert f"{CARPETA}/Firmas Actas/2026-00051/cliente.png" in sp.archivos
    assert f"{CARPETA}/Firmas Actas/2026-00051/representante.png" in sp.archivos
    fila = _fila(sp)
    assert fila["N° de Acta"].value == "2026-00051"
    assert _enlace(fila["PDF original"]) == resultado.ubicacion_pdf


def test_numero_repetido_no_se_guarda(repo, acta_completa):
    repo.guardar(acta_completa, b"%PDF", "a.pdf")
    with pytest.raises(ActaDuplicadaError):
        repo.guardar(acta_completa, b"%PDF", "b.pdf")
    assert repo.numeros() == ["2026-00051"]


def test_si_otro_usuario_guarda_al_mismo_tiempo_no_se_pierde_ninguna_acta(repo, sp, acta_completa):
    repo.guardar(acta_completa, b"%PDF", "1.pdf")

    # Mientras este usuario guarda la 00052, otro guarda la 00053.
    otro = RepositorioSharePoint(sp, carpeta=CARPETA, segundos_cache=0)

    def guardar_otra():
        from copy import deepcopy

        otra = deepcopy(acta_completa)
        otra.numero, otra.fecha_registro = "2026-00053", None
        otro.guardar(otra, b"%PDF", "3.pdf")

    sp.antes_de_escribir = guardar_otra
    acta_completa.numero, acta_completa.fecha_registro = "2026-00052", None
    repo.guardar(acta_completa, b"%PDF", "2.pdf")

    assert sorted(repo.numeros()) == ["2026-00051", "2026-00052", "2026-00053"]


def test_corregir_actualiza_la_fila_y_conserva_el_pdf_original(repo, sp, acta_completa):
    acta_completa.articulos = [Articulo("A-1", "Filtro", 2)]
    repo.guardar(acta_completa, b"%PDF orig", "Acta_2026-00051.pdf")

    acta = repo.obtener("2026-00051")
    assert acta.firma_cliente_png  # firmas recuperadas desde SharePoint
    acta.cliente = "Cliente corregido"
    acta.revision, acta.corregido_por, acta.motivo_correccion = 1, "Ana", "Error de tipeo"
    acta.fecha_correccion = datetime(2026, 9, 28, 10, 0)
    resultado = repo.corregir(acta, b"%PDF rev1", "Acta_2026-00051_Rev1.pdf")

    assert resultado.total_actas == 1
    fila = _fila(sp)
    assert fila["Cliente"].value == "Cliente corregido"
    assert _enlace(fila["PDF original"]).endswith("/PDF/Acta_2026-00051.pdf")
    assert _enlace(fila["PDF corregido"]).endswith("/PDF/Acta_2026-00051_Rev1.pdf")
    assert sp.archivos[f"{CARPETA}/PDF/Acta_2026-00051.pdf"][0] == b"%PDF orig"


def test_catalogo_y_enlaces(repo, sp):
    assert repo.leer_equipos() is None
    sp.archivos[f"{CARPETA}/Equipos.xlsx"] = (b"xlsx", 1)
    assert repo.leer_equipos() == b"xlsx"
    assert repo.enlace_carpeta() == f"https://sp.example/{CARPETA}"
    assert repo.exportar_zip() is None


def test_probar_conexion_informa_cada_paso(repo):
    pasos = repo.probar_conexion()
    assert all(ok for ok, _ in pasos[:1])
    assert any("puede escribir" in m for ok, m in pasos if ok)
    assert any("Actas.xlsx aún no existe" in m for _, m in pasos)


# ---------- AlmacenGraph contra un Graph simulado ----------
class Respuesta:
    def __init__(self, status, json=None, content=b"", headers=None):
        self.status_code, self._json, self.content = status, json or {}, content
        self.headers, self.text = headers or {}, str(json or "")

    def json(self):
        return self._json


class GraphFalso:
    def __init__(self):
        self.peticiones = []
        self.respuestas = {}

    def request(self, metodo, url, headers=None, **kwargs):
        self.peticiones.append((metodo, url, headers, kwargs))
        for (m, fragmento), respuesta in self.respuestas.items():
            if m == metodo and fragmento in url:
                return respuesta(kwargs, headers) if callable(respuesta) else respuesta
        return Respuesta(404)


@pytest.fixture
def graph(monkeypatch):
    falso = GraphFalso()
    falso.respuestas[("GET", "/sites/sistemasanaliticospe.sharepoint.com:/sites/OperacionesyServicios")] = \
        Respuesta(200, {"id": "SITIO"})
    falso.respuestas[("GET", "/sites/SITIO/drives")] = Respuesta(200, {"value": [
        {"id": "OTRA", "name": "Activos del sitio", "webUrl": "https://x/sites/OperacionesyServicios/SiteAssets"},
        {"id": "DOCS", "name": "Documentos",
         "webUrl": "https://x/sites/OperacionesyServicios/Documentos%20compartidos"},
    ]})
    almacen = AlmacenGraph("t", "c", "s", sesion=falso)
    monkeypatch.setattr(almacen, "_token", lambda: "TOKEN")
    return falso, almacen


def test_graph_encuentra_la_biblioteca_por_su_url_y_codifica_la_ruta(graph):
    falso, almacen = graph
    falso.respuestas[("GET", "/drives/DOCS/root:/16.%20Analisis%20de%20Datos/Actas/Actas.xlsx")] = \
        Respuesta(200, {"id": "ITEM", "eTag": '"e1"', "webUrl": "https://x/Actas.xlsx"})
    falso.respuestas[("GET", "/drives/DOCS/items/ITEM/content")] = Respuesta(200, content=b"XLSX")

    archivo = almacen.leer(f"{CARPETA}/Actas.xlsx")
    assert (archivo.datos, archivo.version, archivo.enlace) == (b"XLSX", '"e1"', "https://x/Actas.xlsx")
    assert all(h["Authorization"] == "Bearer TOKEN" for _, _, h, _ in falso.peticiones)


def test_graph_escribe_con_if_match_y_detecta_conflictos(graph):
    falso, almacen = graph
    falso.respuestas[("PUT", "/root:/16.%20Analisis%20de%20Datos/Actas/Actas.xlsx:/content")] = \
        lambda kw, h: Respuesta(412) if h.get("If-Match") == '"viejo"' else Respuesta(200, {"webUrl": "https://x/a"})

    assert almacen.escribir(f"{CARPETA}/Actas.xlsx", b"x", version_esperada='"nuevo"') == "https://x/a"
    with pytest.raises(ConflictoDeVersion):
        almacen.escribir(f"{CARPETA}/Actas.xlsx", b"x", version_esperada='"viejo"')


def test_graph_sin_permiso_en_el_sitio_explica_que_falta(graph):
    falso, almacen = graph
    falso.respuestas[("PUT", "/content")] = Respuesta(403, {"error": "accessDenied"})
    with pytest.raises(AlmacenamientoError, match="Sites.Selected"):
        almacen.escribir(f"{CARPETA}/PDF/a.pdf", b"x")


def test_graph_archivo_inexistente_devuelve_none(graph):
    _, almacen = graph
    assert almacen.leer(f"{CARPETA}/Equipos.xlsx") is None
    assert almacen.enlace(f"{CARPETA}/nada") is None


def test_tenant_invalido_da_un_error_claro_y_no_rompe_la_app():
    almacen = AlmacenGraph("tenant-que-no-existe", "c", "s", sesion=GraphFalso())
    with pytest.raises(AlmacenamientoError, match="tenant_id"):
        almacen.leer(f"{CARPETA}/Actas.xlsx")


def _enlace(celda):
    """URL de la fórmula =HYPERLINK que la app escribe en las columnas de PDF."""
    return leer_hipervinculo(celda.value)[0]


def test_equipo_nuevo_crea_equipos_nuevos_con_el_primero_y_no_duplica(repo, sp, acta_completa):
    from acta_app.catalogo import Catalogo
    from acta_app.equipos_nuevos import es_equipo_nuevo

    ruta = f"{CARPETA}/Equipos_nuevos.xlsx"
    assert repo.leer_equipos_nuevos() is None  # no existe hasta el primer equipo nuevo
    assert repo.registrar_equipo_nuevo(acta_completa, "ana@sistemasanaliticos.com")
    assert ruta in sp.archivos
    assert not repo.registrar_equipo_nuevo(acta_completa, "otro")  # misma serie: no se repite

    otra = acta_completa
    otra.numero, otra.numero_serie, otra.equipo = "2026-00052", "SN-999", "Centrífuga"
    assert repo.registrar_equipo_nuevo(otra, "ana@sistemasanaliticos.com")

    ws = load_workbook(io.BytesIO(sp.archivos[ruta][0])).active
    filas = list(ws.iter_rows(values_only=True))
    assert filas[0][:6] == ("Descripcion", "Marca", "Modelo", "Serie", "Sedes", "Departamentos")
    assert [f[3] for f in filas[1:]] == ["SN123", "SN-999"]
    assert filas[2][6] == "2026-00052" and filas[2][8] == "ana@sistemasanaliticos.com"
    assert ws.tables["TablaEquiposNuevos"].ref == "A1:J3"

    # El autocompletado ya los conoce: deja de ser "nuevo".
    catalogo = Catalogo.desde_bytes(repo.leer_equipos_nuevos())
    assert catalogo.opciones("equipo", {}) and not es_equipo_nuevo(catalogo, otra)


def test_nombres_de_ingenieros_desde_un_excel(repo, sp):
    from openpyxl import Workbook

    wb = Workbook()
    for fila in [["N°", "Nombres y apellidos"], [1, "Ana Ruiz"], [2, "  luis  pérez "], [3, "Ana Ruiz"]]:
        wb.active.append(fila)
    salida = io.BytesIO()
    wb.save(salida)
    sp.escribir(f"{CARPETA}/Firmas Ingenieros/Nombres Ingenieria/Ingenieros.xlsx", salida.getvalue())

    nombres, origen = repo.leer_nombres_ingenieros()
    assert nombres == ["Ana Ruiz", "luis pérez"]
    assert origen.endswith("Nombres Ingenieria/Ingenieros.xlsx")


def test_nombres_de_ingenieros_desde_subcarpetas_junto_a_actas(repo, sp):
    base = "16. Analisis de Datos/Firmas Ingenieros/Nombres Ingenieria"
    sp.carpetas |= {base, f"{base}/Carlos Díaz", f"{base}/Ana Ruiz"}
    sp.escribir(f"{base}/Ana Ruiz/firma.png", b"png")

    nombres, origen = repo.leer_nombres_ingenieros()
    assert nombres == ["Ana Ruiz", "Carlos Díaz"]
    assert origen == f"subcarpetas de {base}"


def test_sin_lista_de_ingenieros(repo):
    assert repo.leer_nombres_ingenieros() is None


def test_nombres_de_ingenieros_desde_el_excel_suelto_en_firmas_ingenieros(repo, sp):
    from openpyxl import Workbook

    wb = Workbook()
    for fila in [["NOMBRES"], ["Sebastián Campos"], ["Ana Ruiz"]]:
        wb.active.append(fila)
    salida = io.BytesIO()
    wb.save(salida)
    sp.escribir(f"{CARPETA}/Firmas Ingenieros/Nombres Ingenieria.xlsx", salida.getvalue())

    nombres, origen = repo.leer_nombres_ingenieros()
    assert nombres == ["Ana Ruiz", "Sebastián Campos"]
    assert origen == f"{CARPETA}/Firmas Ingenieros/Nombres Ingenieria.xlsx"


def test_encuesta_se_guarda_en_la_fila_y_crea_sus_columnas(repo, sp, acta_completa):
    from acta_app.models import EncuestaSatisfaccion
    from acta_app.storage import EncuestaYaRespondidaError

    repo.guardar(acta_completa, b"%PDF", "Acta_2026-00051.pdf")
    assert "Puntualidad del trabajador" not in _fila(sp)  # sin encuestas, no hay columnas

    puntajes = dict.fromkeys(
        ["Puntualidad del trabajador", "Respeto y disposición", "Claridad en la explicación técnica",
         "Orden y limpieza al terminar", "Eficiencia en el trabajo"], 5)
    repo.guardar_encuesta("2026-00051", EncuestaSatisfaccion(puntajes, "Muy buen servicio"))

    ws = load_workbook(io.BytesIO(sp.archivos[f"{CARPETA}/Actas.xlsx"][0]))["Actas"]
    encabezados = [c.value for c in ws[2]]
    inicio = encabezados.index("Puntualidad del trabajador")
    assert encabezados[inicio - 1] == "PDF corregido"
    assert ws.cell(1, inicio + 1).value == "Encuesta de satisfacción del servicio"
    fila = _fila(sp)
    assert fila["Calificación final (/20)"].value == 20
    assert fila["Comentarios y sugerencias"].value == "Muy buen servicio"
    assert fila["Fecha de la encuesta"].value is not None

    with pytest.raises(EncuestaYaRespondidaError):
        repo.guardar_encuesta("2026-00051", EncuestaSatisfaccion(puntajes))

    # Una corrección del acta conserva la encuesta.
    acta = repo.obtener("2026-00051")
    assert acta.encuesta.nota == 20
    acta.encuesta = None
    acta.revision, acta.corregido_por, acta.motivo_correccion = 1, "Ana", "Tipeo"
    repo.corregir(acta, b"%PDF rev1", "Acta_2026-00051_Rev1.pdf")
    assert _fila(sp)["Calificación final (/20)"].value == 20


def test_nota_de_la_encuesta_sobre_20():
    from acta_app.models import EncuestaSatisfaccion

    aspectos = ["Puntualidad del trabajador", "Respeto y disposición", "Claridad en la explicación técnica",
                "Orden y limpieza al terminar", "Eficiencia en el trabajo"]
    assert EncuestaSatisfaccion(dict.fromkeys(aspectos, 5)).nota == 20
    assert EncuestaSatisfaccion(dict.fromkeys(aspectos, 1)).nota == 4
    assert EncuestaSatisfaccion(dict(zip(aspectos, [5, 4, 4, 3, 5]))).nota == 16.8


# ---------- QR de la encuesta ----------
def _codigo_del_enlace(enlace):
    from urllib.parse import parse_qs, urlparse

    return parse_qs(urlparse(enlace).query)["t"][0]


def _puntajes():
    return dict.fromkeys(["Puntualidad del trabajador", "Respeto y disposición",
                          "Claridad en la explicación técnica", "Orden y limpieza al terminar",
                          "Eficiencia en el trabajo"], 4)


def test_qr_de_la_encuesta_de_un_solo_uso(repo, sp, acta_completa):
    from acta_app.encuesta_qr import generar_qr, hash_codigo
    from acta_app.models import EncuestaSatisfaccion
    from acta_app.storage import EnlaceEncuestaInvalidoError, EncuestaYaRespondidaError

    repo.guardar(acta_completa, b"%PDF", "Acta_2026-00051.pdf")
    assert "QR de la encuesta generado el" not in _fila(sp)  # sin QR, sin columnas de encuesta
    acceso, enlace = generar_qr(repo, "2026-00051", "https://sisa-app.streamlit.app/")
    codigo = _codigo_del_enlace(enlace)
    assert enlace.startswith("https://sisa-app.streamlit.app/?encuesta=2026-00051&t=")
    assert (acceso.vence - acceso.generado).total_seconds() == 24 * 3600

    fila = _fila(sp)
    assert fila["Código de acceso (hash)"].value == hash_codigo(codigo)
    assert fila["QR de la encuesta vence el"].value is not None
    assert codigo not in str([c.value for c in fila.values()])  # el código no se guarda

    with pytest.raises(EnlaceEncuestaInvalidoError):
        repo.guardar_encuesta("2026-00051", EncuestaSatisfaccion(_puntajes()), hash_codigo("otro"))
    repo.guardar_encuesta("2026-00051", EncuestaSatisfaccion(_puntajes()), hash_codigo(codigo))
    assert _fila(sp)["Calificación final (/20)"].value == 16
    assert _fila(sp)["Código de acceso (hash)"].value == hash_codigo(codigo)  # se conserva
    # Ya usado: ni el mismo QR ni uno nuevo sirven.
    with pytest.raises(EncuestaYaRespondidaError):
        repo.guardar_encuesta("2026-00051", EncuestaSatisfaccion(_puntajes()), hash_codigo(codigo))
    with pytest.raises(EncuestaYaRespondidaError):
        generar_qr(repo, "2026-00051", "https://a")


def test_qr_nuevo_anula_el_anterior_y_el_qr_vence(repo, sp, acta_completa):
    from datetime import timedelta

    from acta_app.encuesta_qr import generar_qr, hash_codigo
    from acta_app.models import EncuestaSatisfaccion
    from acta_app.storage import EnlaceEncuestaInvalidoError

    repo.guardar(acta_completa, b"%PDF", "Acta_2026-00051.pdf")
    _, enlace_viejo = generar_qr(repo, "2026-00051", "https://a")
    generar_qr(repo, "2026-00051", "https://a")

    with pytest.raises(EnlaceEncuestaInvalidoError):
        repo.guardar_encuesta("2026-00051", EncuestaSatisfaccion(_puntajes()),
                              hash_codigo(_codigo_del_enlace(enlace_viejo)))

    acta = repo.obtener("2026-00051")
    assert acta.acceso_encuesta.vigente()
    assert not acta.acceso_encuesta.vigente(acta.acceso_encuesta.vence + timedelta(minutes=1))


def test_imagen_del_qr_es_un_png():
    from acta_app.encuesta_qr import imagen_qr

    png = imagen_qr("https://sisa-app.streamlit.app/?encuesta=2026-00051&t=" + "x" * 43)
    assert png.startswith(b"\x89PNG")
