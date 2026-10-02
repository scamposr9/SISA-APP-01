import io

from openpyxl import Workbook

from acta_app.protocolos import Protocolos


def _bytes(wb: Workbook) -> bytes:
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def test_formato_en_bloques_ignora_serie_clientes_cronograma_y_obs():
    wb = Workbook()
    ws = wb.active
    ws.title = "Centrífugas"
    filas = [
        ["PROTOCOLO DE MANTENIMIENTO PREVENTIVO"],
        ["CLIENTE:", "Hospital Rebagliati", None, "SEDE:", "Lima"],
        ["EQUIPO:", "Centrífuga", None, "N° SERIE:", "SN-1"],
        ["MARCA:", "Hettich", None, "MODELO:", "Rotina 380"],
        ["N°", "PARTE MANTENIDA", "CRONOGRAMA 2026", "ENE", "FEB"],
        [1, "1. Limpieza general del equipo", None, "X", None],
        [2, "Revisión de escobillas", None, None, "X"],
        [None, None],
        ["OBS:", "Cliente solicita visita en marzo"],
        [None, "1-"],
        [None, "2-"],
        [None, "Texto bajo OBS que no es una actividad"],
        ["EQUIPO: Analizador"],
        ["MARCA: Abbott", None, None, "MODELO: Architect c4000; C8000"],
        ["N°", "Parte mantenida", "Cronograma 2026"],
        [1, "Calibración de pipetas", "X"],
        [2, "Cambio de filtros", None],
    ]
    for fila in filas:
        ws.append(fila)

    protocolos = Protocolos.desde_bytes(_bytes(wb))

    centrifuga = protocolos.buscar("", "HETTICH", "rotina 380")
    assert centrifuga.equipo == "Centrífuga"
    assert centrifuga.actividades == ["Limpieza general del equipo", "Revisión de escobillas"]
    analizador = protocolos.buscar("", "Abbott", "C8000")
    assert analizador.actividades == ["Calibración de pipetas", "Cambio de filtros"]
    assert protocolos.buscar("", "Abbott", "Otro modelo") is None
    assert len(protocolos.lista) == 2


def test_formato_tabla_con_celdas_combinadas_vacias():
    wb = Workbook()
    ws = wb.active
    for fila in [
        ["Equipo", "Marca", "Modelo", "N° Serie", "Parte mantenida", "Cronograma 2026", "OBS"],
        ["Autoclave", "Tuttnauer", "3870EA", "S1", "Revisión de empaquetadura", "X", "ojo"],
        [None, None, None, None, "Limpieza de cámara", "", ""],
        ["Autoclave", "Tuttnauer", "2540M", "S2", "Prueba de presión", "X", ""],
    ]:
        ws.append(fila)

    protocolos = Protocolos.desde_bytes(_bytes(wb))

    assert protocolos.buscar("Autoclave", "Tuttnauer", "3870EA").actividades == [
        "Revisión de empaquetadura", "Limpieza de cámara",
    ]
    assert protocolos.buscar("", "tuttnauer", "2540m").actividades == ["Prueba de presión"]


def test_modelo_unico_se_acepta_aunque_la_marca_no_coincida():
    wb = Workbook()
    ws = wb.active
    for fila in [["MARCA:", "Mindray"], ["MODELO:", "BC-5150"], ["Parte mantenida"], ["Limpieza"]]:
        ws.append(fila)
    protocolos = Protocolos.desde_bytes(_bytes(wb))
    assert protocolos.buscar("", "", "BC-5150").actividades == ["Limpieza"]


def test_renglones_sin_texto_no_son_actividades():
    wb = Workbook()
    ws = wb.active
    for fila in [["MARCA:", "Labtech"], ["MODELO:", "Auto Elisa PW"], ["Parte mantenida"],
                 ["Limpieza exterior"], ["1-"], ["2."], ["-"], ["Limpieza interior"]]:
        ws.append(fila)
    protocolo = Protocolos.desde_bytes(_bytes(wb)).buscar("", "LABTECH", "AUTO ELISA PW")
    assert protocolo.actividades == ["Limpieza exterior", "Limpieza interior"]


def test_mismo_equipo_marca_y_modelo_se_unen_sin_repetir_actividades():
    wb = Workbook()
    ws = wb.active
    for fila in [
        ["EQUIPO:", "Lavador de microplacas"], ["MARCA:", "Labtech"], ["MODELO:", "Auto Elisa PW"],
        ["Parte mantenida"], ["Limpieza exterior"], ["Revisión de bomba"],
        ["OBS:"], ["1-"],
        ["EQUIPO:", "Lavador de microplacas"], ["MARCA:", "LABTECH"], ["MODELO:", "auto elisa pw"],
        ["Parte mantenida"], ["Limpieza exterior"], ["Limpieza de peines de lavado"],
        ["EQUIPO:", "Lector de microplacas"], ["MARCA:", "Labtech"], ["MODELO:", "Auto Elisa PW"],
        ["Parte mantenida"], ["Calibración de filtros"],
    ]:
        ws.append(fila)

    protocolos = Protocolos.desde_bytes(_bytes(wb))

    lavador = protocolos.buscar("Lavador de microplacas", "Labtech", "Auto Elisa PW")
    assert lavador.actividades == ["Limpieza exterior", "Revisión de bomba", "Limpieza de peines de lavado"]
    assert lavador.repetido and lavador.registros == 2
    # Otro equipo con la misma marca y modelo no se mezcla.
    lector = protocolos.buscar("Lector de microplacas", "Labtech", "Auto Elisa PW")
    assert lector.actividades == ["Calibración de filtros"] and not lector.repetido


def test_hojas_que_no_dan_protocolo_quedan_con_su_motivo():
    wb = Workbook()
    wb.active.title = "Buena"
    for fila in [["MARCA:", "Hettich"], ["MODELO:", "EBA 200"], ["Parte mantenida"], ["Limpieza"]]:
        wb.active.append(fila)
    for nombre, filas in [
        ("Sin parte", [["MARCA:", "X"], ["MODELO:", "Y"], ["Actividad"], ["Limpieza"]]),
        ("Sin marca", [["Equipo de prueba"], ["Parte mantenida"], ["Limpieza"]]),
        ("Vacía", []),
    ]:
        ws = wb.create_sheet(nombre)
        for fila in filas:
            ws.append(fila)

    protocolos = Protocolos.desde_bytes(_bytes(wb))

    assert [p.hojas for p in protocolos.lista] == [["Buena"]]
    assert dict(protocolos.descartes) == {
        "Sin parte": "no tiene la columna «Parte mantenida»",
        "Sin marca": "no se encontró MARCA ni MODELO",
        "Vacía": "hoja vacía",
    }
