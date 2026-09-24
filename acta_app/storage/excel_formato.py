"""Leer y escribir el libro del Excel maestro a partir de `Registro`s.

Es independiente de dónde esté el archivo: el repositorio local lo usa sobre disco y el
de SharePoint lo usará sobre el archivo descargado de la biblioteca.

Hoja "Actas":
    fila 1  -> encabezados combinados de cada grupo ("Antecedentes Iniciales", ...)
    fila 2  -> encabezado de cada columna ("Antecedente 1", "Antecedente 2", ...)
    fila 3+ -> una fila por acta
Hoja "Artículos": una fila por artículo empleado, enlazada por "N° de Acta".
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from acta_app import config
from acta_app.storage.esquema import (
    COLUMNA_PDF,
    FORMATO_FECHA,
    GRUPOS,
    Campo,
    Columna,
    Registro,
    disposicion,
    interpretar_encabezados,
)

HOJA_ACTAS = "Actas"
HOJA_ARTICULOS = "Artículos"
TABLA_ACTAS = "TablaActas"
TABLA_ARTICULOS = "TablaArticulos"
FILA_GRUPOS, FILA_ENCABEZADO = 1, 2

COLUMNAS_ARTICULOS = ["N° de Acta", "Fecha", "Cliente", "Código", "Descripción", "Cantidad"]
ANCHOS_ARTICULOS = [14, 12, 30, 16, 45, 10]

_FILL = PatternFill("solid", start_color=config.NAVY.lstrip("#"))
_FILL_GRUPO = PatternFill("solid", start_color=config.NAVY_DARK.lstrip("#"))
_FONT = Font(bold=True, color="FFFFFF")
_BORDE = Border(left=Side(style="thin", color="FFFFFF"), right=Side(style="thin", color="FFFFFF"))
_FONT_ENLACE = Font(color="0563C1", underline="single")


def es_formato_anterior(ws: Worksheet) -> bool:
    """Versión 0.5: encabezados en la fila 1 y textos unidos con ' | '."""
    return ws.cell(FILA_GRUPOS, 1).value == "N° de Acta"


# ---------- Lectura ----------
def leer_registros(ws: Worksheet) -> list[Registro]:
    encabezados = [c.value for c in ws[FILA_ENCABEZADO]]
    columnas = interpretar_encabezados(encabezados)
    registros = []
    for fila in ws.iter_rows(min_row=FILA_ENCABEZADO + 1):
        if all(c.value in (None, "") for c in fila):
            continue
        registro = Registro()
        items: dict[str, dict[int, list]] = {g.titulo: {} for g in GRUPOS}
        for celda, columna in zip(fila, columnas):
            if columna is None:
                continue
            if isinstance(columna.bloque, Campo):
                registro.valores[columna.bloque.nombre] = celda.value
                if columna.bloque.nombre == COLUMNA_PDF and celda.hyperlink is not None:
                    registro.enlace_pdf = celda.hyperlink.target
            else:
                item = items[columna.bloque.titulo].setdefault(
                    columna.item, [None] * columna.bloque.columnas_por_item
                )
                item[columna.subcampo] = celda.value
        for titulo, por_indice in items.items():
            registro.items[titulo] = [
                tuple(valores)
                for _, valores in sorted(por_indice.items())
                if any(v not in (None, "") for v in valores)
            ]
        registros.append(registro)
    return registros


# ---------- Escritura ----------
def construir_libro(registros: list[Registro]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA_ACTAS
    columnas = disposicion(registros)
    _encabezados(ws, columnas)
    for n, registro in enumerate(registros, start=FILA_ENCABEZADO + 1):
        for i, columna in enumerate(columnas, start=1):
            celda = ws.cell(n, i, columna.valor(registro))
            celda.alignment = Alignment(vertical="top", wrap_text=columna.texto_largo)
            if columna.formato:
                celda.number_format = columna.formato
            if columna.encabezado == COLUMNA_PDF and registro.enlace_pdf:
                celda.hyperlink = registro.enlace_pdf
                celda.font = _FONT_ENLACE
    _tabla(ws, TABLA_ACTAS, len(columnas), FILA_ENCABEZADO, len(registros))
    ws.freeze_panes = ws.cell(FILA_ENCABEZADO + 1, 2)

    _hoja_articulos(wb.create_sheet(HOJA_ARTICULOS), registros)
    return wb


def _encabezados(ws: Worksheet, columnas: list[Columna]) -> None:
    for i, columna in enumerate(columnas, start=1):
        for fila, fill in ((FILA_GRUPOS, _FILL_GRUPO), (FILA_ENCABEZADO, _FILL)):
            celda = ws.cell(fila, i)
            celda.fill, celda.font, celda.border = fill, _FONT, _BORDE
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.cell(FILA_ENCABEZADO, i, columna.encabezado)
        ws.column_dimensions[get_column_letter(i)].width = columna.ancho

    # Encabezado combinado sobre las columnas de cada grupo.
    for grupo in GRUPOS:
        posiciones = [i for i, c in enumerate(columnas, start=1) if c.bloque is grupo]
        ws.cell(FILA_GRUPOS, posiciones[0], grupo.titulo)
        if len(posiciones) > 1:
            ws.merge_cells(
                start_row=FILA_GRUPOS, start_column=posiciones[0],
                end_row=FILA_GRUPOS, end_column=posiciones[-1],
            )
    ws.row_dimensions[FILA_GRUPOS].height = 20
    ws.row_dimensions[FILA_ENCABEZADO].height = 30


def _hoja_articulos(ws: Worksheet, registros: list[Registro]) -> None:
    ws.append(COLUMNAS_ARTICULOS)
    for i, (nombre, ancho) in enumerate(zip(COLUMNAS_ARTICULOS, ANCHOS_ARTICULOS), start=1):
        celda = ws.cell(1, i)
        celda.fill, celda.font = _FILL, _FONT
        ws.column_dimensions[get_column_letter(i)].width = ancho
    filas = 0
    for registro in registros:
        for codigo, descripcion, cantidad in registro.items.get("Artículos Empleados", []):
            v = registro.valores
            ws.append([v.get("N° de Acta"), v.get("Fecha"), v.get("Cliente"), codigo, descripcion, cantidad])
            ws.cell(ws.max_row, 2).number_format = FORMATO_FECHA
            filas += 1
    ws.freeze_panes = "B2"
    if filas:
        _tabla(ws, TABLA_ARTICULOS, len(COLUMNAS_ARTICULOS), 1, filas)


def _tabla(ws: Worksheet, nombre: str, n_columnas: int, fila_encabezado: int, n_filas: int) -> None:
    """Tabla de Excel (filtros, franjas; y es la base para la API de Excel de SharePoint)."""
    ultima = fila_encabezado + max(n_filas, 1)
    tabla = Table(displayName=nombre, ref=f"A{fila_encabezado}:{get_column_letter(n_columnas)}{ultima}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(tabla)
