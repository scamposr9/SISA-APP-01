from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from acta_app.models import Articulo
from acta_app.storage import (
    COLUMNAS_ACTAS,
    COLUMNAS_ARTICULOS,
    ActaDuplicadaError,
    RepositorioExcelLocal,
)


@pytest.fixture
def repo(tmp_path):
    return RepositorioExcelLocal(tmp_path / "actas_maestro.xlsx", tmp_path / "pdfs")


def test_primera_acta_crea_el_excel_con_encabezados_y_una_fila(repo, acta_completa):
    resultado = repo.guardar(acta_completa, b"%PDF-1.4", "Acta_2026-00051.pdf")

    assert resultado.total_actas == 1
    assert (repo.dir_pdf / "Acta_2026-00051.pdf").read_bytes() == b"%PDF-1.4"
    wb = load_workbook(repo.ruta_excel)
    ws = wb["Actas"]
    assert [c.value for c in ws[1]] == COLUMNAS_ACTAS
    fila = dict(zip(COLUMNAS_ACTAS, [c.value for c in ws[2]]))
    assert fila["N° de Acta"] == "2026-00051"
    assert fila["Cliente"] == "Hospital Rebagliati"
    assert fila["Fecha"] == datetime(2026, 9, 24)  # fecha real de Excel, no texto
    assert isinstance(fila["Fecha de registro"], datetime)
    assert fila["Archivo PDF"] == "Acta_2026-00051.pdf"
    assert ws.tables["TablaActas"].ref == f"A1:{ws.cell(1, len(COLUMNAS_ACTAS)).column_letter}2"


def test_articulos_van_a_su_propia_hoja_una_fila_por_articulo(repo, acta_completa):
    acta_completa.articulos = [Articulo("A-1", "Filtro", 2), Articulo(), Articulo("B-2", "Sello", 1)]
    repo.guardar(acta_completa, b"pdf", "a.pdf")

    ws = load_workbook(repo.ruta_excel)["Artículos"]
    assert [c.value for c in ws[1]] == COLUMNAS_ARTICULOS
    filas = [[c.value for c in fila] for fila in ws.iter_rows(min_row=2)]
    assert [(f[0], f[3], f[5]) for f in filas] == [("2026-00051", "A-1", 2), ("2026-00051", "B-2", 1)]


def test_varias_actas_se_agregan_como_filas_nuevas(repo, acta_completa):
    for numero in ["2026-00051", "2026-00052", "2026-00053"]:
        acta_completa.numero = numero
        acta_completa.fecha_registro = None
        repo.guardar(acta_completa, b"pdf", f"{numero}.pdf")

    df = repo.leer_actas()
    assert list(df["N° de Acta"]) == ["2026-00051", "2026-00052", "2026-00053"]


def test_numero_repetido_no_se_guarda(repo, acta_completa):
    repo.guardar(acta_completa, b"pdf", "a.pdf")
    acta_completa.numero = " 2026-00051 "  # mismo número con espacios
    assert repo.existe(acta_completa.numero)
    with pytest.raises(ActaDuplicadaError):
        repo.guardar(acta_completa, b"pdf", "b.pdf")
    assert len(repo.leer_actas()) == 1


def test_sin_excel_no_hay_actas(repo):
    assert not repo.existe("2026-00001")
    assert repo.excel_bytes() is None
    assert repo.leer_actas().empty
