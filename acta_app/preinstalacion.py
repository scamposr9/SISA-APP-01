"""Reporte de preinstalación (Presite): datos, validación y su propio Excel maestro.

La preinstalación revisa aspectos que no se ven en un mantenimiento (condiciones eléctricas,
traslado, accesos, área, complementos faltantes, contactos), así que se guarda en un Excel
aparte (Preinstalaciones.xlsx) para no cambiar la estructura de Actas.xlsx.

Excel: una fila por reporte. Las listas (accesos, contactos, observaciones) ocupan tantas
columnas como el reporte que más tenga («Acceso 1», «Acceso 2», …), igual que en Actas.xlsx.
Las opciones marcadas van como «X» (complementos) o con su nombre (tomas, traslado, servicio).
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from acta_app import config
from acta_app.catalogo import clave

# ---------- Opciones del formato (mismo orden que el Word) ----------
# Tipos de toma eléctrica: los dibujos del Word, en assets/tomas/toma_NN.png.
TIPOS_TOMA = [f"Tipo {i}" for i in range(1, 14)]
TRASLADO_ESTIBADORES = "Estibadores"
TRASLADO_NO_REQUIERE = "No requiere"
TRASLADOS = ["Estoca", "Apilador", TRASLADO_ESTIBADORES, TRASLADO_NO_REQUIERE]
SERVICIO_LABORATORIO = "Laboratorio"
SERVICIOS_AREA = ["Banco de Sangre", SERVICIO_LABORATORIO]
COMPLEMENTOS = ["Aire Acondicionado", "Lavaderos", "Punto de Agua", "Punto de Desagüe", "Puntos de Red", "Calefacción"]
MEDIDAS = ["Largo", "Ancho", "Altura"]
SUPERFICIES = ["Mesa de Trabajo", "Piso"]
ACCESOS_INICIALES = 3  # filas de «Accesos» que aparecen de entrada


def imagen_toma(tipo: str) -> bytes:
    numero = TIPOS_TOMA.index(tipo) + 1
    return (config.ASSETS_DIR / "tomas" / f"toma_{numero:02d}.png").read_bytes()


@dataclass
class Contacto:
    nombre: str = ""
    cargo: str = ""
    telefono: str = ""

    @property
    def esta_vacio(self) -> bool:
        return not (self.nombre or self.cargo or self.telefono)


@dataclass
class Preinstalacion:
    numero: str = ""
    fecha: date | None = None
    cliente: str = ""
    ubicacion: str = ""
    equipo: str = ""
    marca: str = ""
    modelo: str = ""
    numero_serie: str = ""

    # Condiciones eléctricas
    punto_dedicado: bool | None = None
    tipos_toma: list[str] = field(default_factory=list)

    # Detalles
    traslado: list[str] = field(default_factory=list)
    estibadores: int | None = None  # cuántos, si se marcó «Estibadores»
    accesos: list[str] = field(default_factory=list)

    # Tipo de área
    servicios: list[str] = field(default_factory=list)
    tipo_laboratorio: str = ""

    # Condiciones del área, en cm: {"Mesa de Trabajo": {"Largo": 330, ...}, "Piso": {...}}
    medidas: dict[str, dict[str, float | None]] = field(default_factory=dict)

    complementos_faltantes: list[str] = field(default_factory=list)
    temperatura: str = ""
    contactos: list[Contacto] = field(default_factory=list)
    observaciones: list[str] = field(default_factory=list)

    realizado_por: str = ""
    fecha_registro: datetime | None = None
    registrado_por: str = ""  # cuenta con la que se inició sesión

    @property
    def contactos_usados(self) -> list[Contacto]:
        return [c for c in self.contactos if not c.esta_vacio]

    def medida(self, superficie: str, medida: str) -> float | None:
        return (self.medidas.get(superficie) or {}).get(medida)

    def traslado_texto(self, opcion: str) -> str:
        """'Estibadores' -> 'Estibadores (3)' si se indicó cuántos."""
        if opcion == TRASLADO_ESTIBADORES and self.estibadores:
            return f"{opcion} ({self.estibadores})"
        return opcion


def a_dict(p: Preinstalacion) -> dict:
    """Apartados propios de la preinstalación, para el borrador (JSON)."""
    return {
        "punto_dedicado": p.punto_dedicado, "tipos_toma": p.tipos_toma, "traslado": p.traslado,
        "estibadores": p.estibadores, "accesos": p.accesos, "servicios": p.servicios,
        "tipo_laboratorio": p.tipo_laboratorio, "medidas": p.medidas,
        "complementos_faltantes": p.complementos_faltantes, "temperatura": p.temperatura,
        "contactos": [[c.nombre, c.cargo, c.telefono] for c in p.contactos_usados],
        "observaciones": p.observaciones, "realizado_por": p.realizado_por,
    }


def desde_dict(d: dict) -> Preinstalacion:
    return Preinstalacion(
        punto_dedicado=d.get("punto_dedicado"),
        tipos_toma=[t for t in d.get("tipos_toma") or [] if t in TIPOS_TOMA],
        traslado=list(d.get("traslado") or []), estibadores=d.get("estibadores"),
        accesos=list(d.get("accesos") or []), servicios=list(d.get("servicios") or []),
        tipo_laboratorio=d.get("tipo_laboratorio") or "", medidas=dict(d.get("medidas") or {}),
        complementos_faltantes=list(d.get("complementos_faltantes") or []),
        temperatura=d.get("temperatura") or "",
        contactos=[Contacto(*c) for c in d.get("contactos") or []],
        observaciones=list(d.get("observaciones") or []), realizado_por=d.get("realizado_por") or "",
    )


def tiene_datos(p: Preinstalacion) -> bool:
    d = a_dict(p)
    d["medidas"] = [v for m in p.medidas.values() for v in m.values() if v is not None]
    return any(v not in (None, "", [], {}) for v in d.values())


def validar(p: Preinstalacion) -> list[str]:
    """Campos faltantes o inválidos (vacío = se puede guardar)."""
    errores = [nombre for nombre, valor in (
        ("N.° de reporte", p.numero), ("Fecha", p.fecha), ("Cliente", p.cliente), ("Ubicación", p.ubicacion),
        ("Equipo", p.equipo), ("Marca", p.marca), ("Modelo", p.modelo),
    ) if not valor]
    if p.punto_dedicado is None:
        errores.append("¿Es punto dedicado?")
    if not p.tipos_toma:
        errores.append("Tipo de toma eléctrica")
    if not p.traslado:
        errores.append("Traslado del equipo")
    elif TRASLADO_NO_REQUIERE in p.traslado and len(p.traslado) > 1:
        errores.append("Traslado del equipo («No requiere» no va con otras opciones)")
    if TRASLADO_ESTIBADORES in p.traslado and not p.estibadores:
        errores.append("Cantidad de estibadores")
    if not p.accesos:
        errores.append("Accesos")
    if not p.servicios:
        errores.append("Tipo de área (servicio)")
    if not p.contactos_usados:
        errores.append("Personal de contacto")
    elif any(not c.nombre or not c.telefono for c in p.contactos_usados):
        errores.append("Personal de contacto (cada contacto necesita nombre y teléfono)")
    if not p.realizado_por:
        errores.append("Realizado por")
    return errores


def nombre_archivo_pdf(p: Preinstalacion) -> str:
    seguro = "".join(ch if (ch.isascii() and ch.isalnum()) or ch == "-" else "_" for ch in p.numero) or "sin_numero"
    return f"Preinstalacion_{seguro}.pdf"


# ---------- Excel maestro ----------
HOJA = "Preinstalaciones"
TABLA = "TablaPreinstalaciones"
COLUMNA_NUMERO = "N° de Reporte"
COLUMNA_PDF = "PDF"
MARCA = "X"
FORMATO_FECHA, FORMATO_FECHA_HORA = "dd/mm/yyyy", "dd/mm/yyyy hh:mm:ss"


def _medida_col(superficie: str, medida: str) -> str:
    return f"{superficie} - {medida} (cm)"


def _falta_col(complemento: str) -> str:
    return f"Falta: {complemento}"


# (encabezado, ancho). "{n}" = columnas repetidas por ítem (Acceso 1, Acceso 2, ...).
_COLUMNAS: list[tuple[str, int]] = [
    (COLUMNA_NUMERO, 14), ("Fecha", 12), ("Cliente", 30), ("Ubicación", 22), ("Equipo", 24), ("Marca", 16),
    ("Modelo", 16), ("N° Serie", 16), ("Punto dedicado", 10), ("Tipo de toma eléctrica", 22), ("Traslado del equipo", 26),
    ("Acceso {n}", 40),
    ("Servicio", 20), ("Tipo de laboratorio", 22),
    *[(_medida_col(s, m), 12) for s in SUPERFICIES for m in MEDIDAS],
    *[(_falta_col(c), 12) for c in COMPLEMENTOS],
    ("Temperatura del área", 22),
    ("Contacto {n} - Nombre", 22), ("Contacto {n} - Cargo", 26), ("Contacto {n} - Teléfono", 16),
    ("Observación {n}", 40),
    ("Realizado por", 26), ("Registrado por", 30), ("Fecha de registro", 20), (COLUMNA_PDF, 28),
]
_GRUPOS = {  # prefijo de columna repetida -> nombres de sus subcolumnas
    "Acceso": ["Acceso {n}"],
    "Contacto": ["Contacto {n} - Nombre", "Contacto {n} - Cargo", "Contacto {n} - Teléfono"],
    "Observación": ["Observación {n}"],
}


def fila(p: Preinstalacion, nombre_pdf: str = "", enlace_pdf: str = "") -> dict[str, object]:
    """Encabezado -> valor de la fila del reporte."""
    v: dict[str, object] = {
        COLUMNA_NUMERO: p.numero, "Fecha": p.fecha, "Cliente": p.cliente, "Ubicación": p.ubicacion,
        "Equipo": p.equipo, "Marca": p.marca, "Modelo": p.modelo, "N° Serie": p.numero_serie,
        "Punto dedicado": None if p.punto_dedicado is None else ("Sí" if p.punto_dedicado else "No"),
        "Tipo de toma eléctrica": ", ".join(p.tipos_toma),
        "Traslado del equipo": ", ".join(p.traslado_texto(t) for t in p.traslado),
        "Servicio": ", ".join(p.servicios),
        "Tipo de laboratorio": p.tipo_laboratorio if SERVICIO_LABORATORIO in p.servicios else "",
        "Temperatura del área": p.temperatura,
        "Realizado por": p.realizado_por, "Registrado por": p.registrado_por,
        "Fecha de registro": p.fecha_registro,
        COLUMNA_PDF: _hipervinculo(enlace_pdf, nombre_pdf) if enlace_pdf else nombre_pdf,
    }
    for s in SUPERFICIES:
        for m in MEDIDAS:
            v[_medida_col(s, m)] = p.medida(s, m)
    for c in COMPLEMENTOS:
        v[_falta_col(c)] = MARCA if c in p.complementos_faltantes else ""
    for n, acceso in enumerate(p.accesos, start=1):
        v[f"Acceso {n}"] = acceso
    for n, contacto in enumerate(p.contactos_usados, start=1):
        v[f"Contacto {n} - Nombre"] = contacto.nombre
        v[f"Contacto {n} - Cargo"] = contacto.cargo
        v[f"Contacto {n} - Teléfono"] = contacto.telefono
    for n, obs in enumerate(p.observaciones, start=1):
        v[f"Observación {n}"] = obs
    return v


def _hipervinculo(enlace: str, texto: str) -> str:
    from acta_app.storage.excel_formato import formula_hipervinculo

    return formula_hipervinculo(enlace, texto) or texto


def leer_filas(contenido: bytes | None) -> list[dict[str, object]]:
    """Filas del Excel como encabezado -> valor (las fórmulas HYPERLINK se conservan)."""
    if not contenido:
        return []
    ws = load_workbook(io.BytesIO(contenido))[HOJA]
    encabezados = [c.value for c in ws[1]]
    filas = []
    for celdas in ws.iter_rows(min_row=2, values_only=True):
        if all(c in (None, "") for c in celdas):
            continue
        filas.append({str(h): v for h, v in zip(encabezados, celdas) if h})
    return filas


def existe(filas: list[dict[str, object]], numero: str) -> bool:
    buscado = clave("".join(numero.split()))
    return any(clave("".join(str(f.get(COLUMNA_NUMERO) or "").split())) == buscado for f in filas)


def _cantidad(filas: list[dict[str, object]], prefijo: str) -> int:
    patron = re.compile(rf"^{re.escape(prefijo)} (\d+)\b")
    numeros = [int(m.group(1)) for f in filas for h, v in f.items() if v not in (None, "") and (m := patron.match(h))]
    return max(numeros + [1])


def encabezados(filas: list[dict[str, object]]) -> list[tuple[str, int]]:
    """Columnas del Excel: cada lista con tantas columnas como el reporte que más tenga."""
    resultado: list[tuple[str, int]] = []
    vistos: set[str] = set()
    for nombre, ancho in _COLUMNAS:
        if "{n}" not in nombre:
            resultado.append((nombre, ancho))
            continue
        prefijo = next(p for p, cols in _GRUPOS.items() if nombre in cols)
        if prefijo in vistos:
            continue
        vistos.add(prefijo)
        anchos = dict(_COLUMNAS)
        for n in range(1, _cantidad(filas, prefijo) + 1):
            resultado += [(col.format(n=n), anchos[col]) for col in _GRUPOS[prefijo]]
    return resultado


def construir_libro(filas: list[dict[str, object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA
    columnas = encabezados(filas)
    relleno, fuente = PatternFill("solid", start_color=config.NAVY.lstrip("#")), Font(bold=True, color="FFFFFF")
    for i, (nombre, ancho) in enumerate(columnas, start=1):
        celda = ws.cell(1, i, nombre)
        celda.fill, celda.font = relleno, fuente
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[1].height = 32
    for n, datos in enumerate(filas, start=2):
        for i, (nombre, ancho) in enumerate(columnas, start=1):
            celda = ws.cell(n, i, datos.get(nombre))
            celda.alignment = Alignment(vertical="top", wrap_text=ancho >= 30)
            if nombre == "Fecha":
                celda.number_format = FORMATO_FECHA
            elif nombre == "Fecha de registro":
                celda.number_format = FORMATO_FECHA_HORA
            elif nombre == COLUMNA_PDF and str(celda.value or "").startswith("="):
                celda.font = Font(color="0563C1", underline="single")
    ultima = 1 + max(len(filas), 1)
    tabla = Table(displayName=TABLA, ref=f"A1:{get_column_letter(len(columnas))}{ultima}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(tabla)
    ws.freeze_panes = "B2"
    wb.calculation.fullCalcOnLoad = True  # calcula los HYPERLINK al abrir
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def agregar(contenido: bytes | None, p: Preinstalacion, nombre_pdf: str, enlace_pdf: str) -> tuple[bytes, int]:
    """(libro con el reporte agregado al final, total de reportes)."""
    filas = leer_filas(contenido)
    filas.append(fila(p, nombre_pdf, enlace_pdf))
    return construir_libro(filas), len(filas)
