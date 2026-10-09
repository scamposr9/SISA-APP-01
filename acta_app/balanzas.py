"""Mantenimiento preventivo de balanzas: «Pruebas de funcionamiento» y su Excel aparte.

El acta de una balanza es igual a las demás (se guarda en Actas.xlsx como siempre), pero
después de «Acciones realizadas» lleva la tabla de pruebas (Equipment Details): por cada peso
de referencia («Required Weight»), lo que mostró la balanza («Displayed Weight») y lo que
muestra tras el ajuste («Adjustment Weight»). La actividad «Pruebas de funcionamiento» del
protocolo no va en el checklist: es esta tabla.

Además, cada acta de balanza se copia a «Mantenimientos Balanzas.xlsx» (una fila por acta:
las mismas columnas de Actas.xlsx más las de las pruebas).
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from acta_app import config
from acta_app.catalogo import clave

# Pesos de referencia (fila «Required Weight»), fijos.
PESOS_REQUERIDOS = [100, 300, 500, 700, 900, 1000, 1290, 1500, 1700, 2000, 2300, 2500]
UNIDAD = "gm"
FILA_REQUERIDO = f"Required Weight ({UNIDAD})"
FILA_MOSTRADO = f"Displayed Weight ({UNIDAD})"
FILA_AJUSTADO = f"Adjustment Weight ({UNIDAD})"
TITULO_TABLA = "Equipment Details"
SECCION = "Pruebas de funcionamiento"


def es_balanza(*nombres: str) -> bool:
    """¿El equipo (o su protocolo) es una balanza? No importan marca ni modelo."""
    return any("balanza" in clave(n) for n in nombres if n)


def es_actividad_de_pruebas(texto: str) -> bool:
    """«Pruebas de funcionamiento» del protocolo: va como tabla, no como casilla."""
    t = clave(texto)
    return "prueba" in t and "funcionamiento" in t


@dataclass
class PruebasBalanza:
    mostrados: list[float | None] = field(default_factory=lambda: [None] * len(PESOS_REQUERIDOS))
    ajustados: list[float | None] = field(default_factory=lambda: [None] * len(PESOS_REQUERIDOS))

    @property
    def tiene_datos(self) -> bool:
        return any(v is not None for v in self.mostrados + self.ajustados)


def validar(p: PruebasBalanza) -> list[str]:
    errores = []
    for fila, valores in ((FILA_MOSTRADO, p.mostrados), (FILA_AJUSTADO, p.ajustados)):
        faltan = [str(peso) for peso, v in zip(PESOS_REQUERIDOS, valores) if v is None]
        if faltan:
            errores.append(f"Pruebas de funcionamiento: {fila} en {', '.join(faltan)}")
    return errores


def numero(v: float | None) -> str:
    """1288.0 -> '1288'; 99.5 -> '99.5'."""
    return "" if v is None else f"{v:g}"


# ---------- Excel «Mantenimientos Balanzas.xlsx» ----------
HOJA = "Balanzas"
TABLA = "TablaBalanzas"
COLUMNA_NUMERO = "N° de Acta"
COLUMNA_PDF = "PDF"


def _columna(fila: str, peso: int) -> str:
    return f"{fila.split(' (')[0]} {peso} {UNIDAD}"  # «Displayed Weight 100 gm»


def columnas_pruebas() -> list[str]:
    return [_columna(f, peso) for f in (FILA_MOSTRADO, FILA_AJUSTADO) for peso in PESOS_REQUERIDOS]


def fila_pruebas(p: PruebasBalanza) -> dict[str, object]:
    valores = {}
    for fila, lista in ((FILA_MOSTRADO, p.mostrados), (FILA_AJUSTADO, p.ajustados)):
        for peso, v in zip(PESOS_REQUERIDOS, lista):
            valores[_columna(fila, peso)] = v
    return valores


def pruebas_de_fila(fila: dict[str, object]) -> PruebasBalanza:
    def leer(nombre: str) -> list[float | None]:
        return [None if fila.get(_columna(nombre, peso)) in (None, "") else float(fila[_columna(nombre, peso)])
                for peso in PESOS_REQUERIDOS]
    return PruebasBalanza(leer(FILA_MOSTRADO), leer(FILA_AJUSTADO))


def leer_filas(contenido: bytes | None) -> list[dict[str, object]]:
    if not contenido:
        return []
    ws = load_workbook(io.BytesIO(contenido))[HOJA]
    encabezados = [c.value for c in ws[1]]
    return [{str(h): v for h, v in zip(encabezados, celdas) if h}
            for celdas in ws.iter_rows(min_row=2, values_only=True) if any(c not in (None, "") for c in celdas)]


def buscar(filas: list[dict[str, object]], numero_acta: str) -> dict[str, object] | None:
    buscado = clave("".join(numero_acta.split()))
    return next((f for f in filas if clave("".join(str(f.get(COLUMNA_NUMERO) or "").split())) == buscado), None)


def registrar(contenido: bytes | None, fila: dict[str, object]) -> tuple[bytes, int]:
    """Agrega la fila del acta o, si ya estaba (corrección), la reemplaza."""
    filas = leer_filas(contenido)
    anterior = buscar(filas, str(fila.get(COLUMNA_NUMERO) or ""))
    if anterior is not None:
        filas[filas.index(anterior)] = fila
    else:
        filas.append(fila)
    return construir_libro(filas), len(filas)


def construir_libro(filas: list[dict[str, object]]) -> bytes:
    # Columnas: las del acta en su orden y, al final, las de las pruebas y el PDF.
    pruebas = columnas_pruebas()
    columnas: list[str] = []
    for f in filas:
        columnas += [c for c in f if c not in columnas and c not in pruebas and c != COLUMNA_PDF]
    columnas += pruebas + [COLUMNA_PDF]

    wb = Workbook()
    ws = wb.active
    ws.title = HOJA
    relleno, fuente = PatternFill("solid", start_color=config.NAVY.lstrip("#")), Font(bold=True, color="FFFFFF")
    for i, nombre in enumerate(columnas, start=1):
        celda = ws.cell(1, i, nombre)
        celda.fill, celda.font = relleno, fuente
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = 12 if nombre in pruebas else 20
    ws.row_dimensions[1].height = 32
    for n, datos in enumerate(filas, start=2):
        for i, nombre in enumerate(columnas, start=1):
            valor = datos.get(nombre)
            celda = ws.cell(n, i, valor)
            celda.alignment = Alignment(vertical="top")
            if nombre == "Fecha":
                celda.number_format = "dd/mm/yyyy"
            elif nombre in ("Fecha de registro", "Fecha de corrección"):
                celda.number_format = "dd/mm/yyyy hh:mm:ss"
            elif nombre in ("Hora Inicio Trabajo", "Hora Fin Trabajo"):
                celda.number_format = "hh:mm"
            elif nombre == COLUMNA_PDF and str(valor or "").startswith("="):
                celda.font = Font(color="0563C1", underline="single")
    tabla = Table(displayName=TABLA, ref=f"A1:{get_column_letter(len(columnas))}{1 + max(len(filas), 1)}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(tabla)
    ws.freeze_panes = "B2"
    wb.calculation.fullCalcOnLoad = True
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
