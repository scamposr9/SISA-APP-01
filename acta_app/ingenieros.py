"""Nombres de los ingenieros que pueden firmar un acta (lista desplegable del representante).

Se leen de la carpeta «Firmas Ingenieros/Nombres Ingenieria» en SharePoint:
  * si contiene un Excel, de su columna «Nombre…», «Ingenieros…», etc. (o de la primera
    columna con texto);
  * si no, de los nombres de sus subcarpetas (una por ingeniero).
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from acta_app.catalogo import clave, limpiar


# Encabezados que indican la columna de nombres («Nombres y apellidos», «Ingenieros», …).
_ENCABEZADOS = ("nombre", "ingeniero", "apellido", "personal", "colaborador", "responsable", "tecnico")


def ordenar(nombres: list[str]) -> list[str]:
    """Sin repetidos (mayúsculas, tildes y espacios no cuentan) y en orden alfabético."""
    unicos: dict[str, str] = {}
    for nombre in map(limpiar, nombres):
        if nombre and clave(nombre) not in unicos:
            unicos[clave(nombre)] = nombre
    return sorted(unicos.values(), key=clave)


def nombres_desde_excel(datos: bytes) -> list[str]:
    wb = load_workbook(io.BytesIO(datos), data_only=True, read_only=True)
    filas = [list(f) for f in wb.worksheets[0].iter_rows(values_only=True)]
    wb.close()
    columna, inicio = None, 0
    for n, fila in enumerate(filas[:10]):
        columna = next((i for i, c in enumerate(fila) if clave(limpiar(c)).startswith(_ENCABEZADOS)), None)
        if columna is not None:
            inicio = n + 1
            break
    if columna is None:  # sin encabezado: primera columna con texto
        columna = next(
            (i for i in range(max((len(f) for f in filas), default=0))
             if any(isinstance(f[i], str) and limpiar(f[i]) for f in filas if i < len(f))),
            0,
        )
    nombres = [f[columna] for f in filas[inicio:] if columna < len(f) and isinstance(f[columna], str)]
    return ordenar(nombres)
