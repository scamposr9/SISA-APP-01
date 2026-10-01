"""Equipos que no están en el catálogo (Equipos.xlsx), anotados para revisarlos.

Cuando se guarda un acta cuyo N.° de serie no está en el catálogo, el equipo se agrega a
«Equipos_nuevos.xlsx» (se crea con el primero). Sus primeras columnas son las mismas que
las de Equipos.xlsx, para copiar las filas revisadas tal cual; las demás indican de qué
acta salió cada equipo. El autocompletado también sugiere estos equipos mientras tanto.
"""

from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from acta_app import config
from acta_app.catalogo import Catalogo, clave, limpiar
from acta_app.models import Acta, ahora

HOJA = "Equipos nuevos"
TABLA = "TablaEquiposNuevos"
# Mismos nombres que en Equipos.xlsx (ver catalogo.COLUMNAS) + datos de origen.
COLUMNAS = [
    ("Descripcion", 30), ("Marca", 20), ("Modelo", 20), ("Serie", 20),
    ("Sedes", 32), ("Departamentos", 20),
    ("N° de Acta", 14), ("Fecha del acta", 14), ("Registrado por", 30), ("Fecha de registro", 20),
]
_COLUMNA_SERIE = 4


def es_equipo_nuevo(catalogo: Catalogo, acta: Acta) -> bool:
    """¿La serie del acta no aparece en el catálogo (Equipos.xlsx + equipos ya anotados)?"""
    serie = clave(acta.numero_serie)
    return bool(serie) and serie not in set(catalogo.datos["serie"].map(clave))


def agregar(contenido: bytes | None, acta: Acta, registrado_por: str) -> bytes | None:
    """Libro con el equipo del acta agregado al final, o None si su serie ya estaba."""
    if contenido:
        wb = load_workbook(io.BytesIO(contenido))
        ws = wb[HOJA] if HOJA in wb.sheetnames else wb.worksheets[0]
    else:
        wb, ws = _libro_nuevo()

    serie = clave(acta.numero_serie)
    for (valor,) in ws.iter_rows(min_row=2, min_col=_COLUMNA_SERIE, max_col=_COLUMNA_SERIE, values_only=True):
        if clave(limpiar(valor)) == serie:
            return None

    ws.append([
        acta.equipo, acta.marca, acta.modelo, acta.numero_serie, acta.cliente, acta.ubicacion,
        acta.numero, acta.fecha, registrado_por, ahora().replace(tzinfo=None),
    ])
    fila = ws.max_row
    ws.cell(fila, 8).number_format = "DD/MM/YYYY"
    ws.cell(fila, 10).number_format = "DD/MM/YYYY HH:MM"
    if TABLA in ws.tables:
        ws.tables[TABLA].ref = f"A1:{get_column_letter(len(COLUMNAS))}{fila}"

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _libro_nuevo():
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA
    relleno = PatternFill("solid", start_color=config.NAVY.lstrip("#"))
    for i, (nombre, ancho) in enumerate(COLUMNAS, start=1):
        celda = ws.cell(1, i, nombre)
        celda.fill, celda.font = relleno, Font(bold=True, color="FFFFFF")
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"
    tabla = Table(displayName=TABLA, ref=f"A1:{get_column_letter(len(COLUMNAS))}2")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tabla)
    return wb, ws
