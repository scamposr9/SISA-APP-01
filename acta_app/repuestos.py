"""Catálogo de repuestos (Repuestos.xlsx en SharePoint) para «Artículos empleados».

Se leen todas las hojas. En cada una se busca la fila de encabezados con una columna de
código y otra de descripción (en las primeras filas); el resto de columnas se ignora.
Los códigos repetidos (en la misma hoja o en otra) se guardan una sola vez, con la
primera descripción encontrada.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field

from openpyxl import load_workbook

from acta_app.catalogo import clave, limpiar

_FILAS_PARA_ENCABEZADO = 20
_COLUMNAS_CODIGO = ("codigo", "cod", "code", "sku", "part number", "numero de parte",
                    "n° de parte", "n° parte", "nro de parte", "p/n", "item")
_COLUMNAS_DESCRIPCION = ("descripcion", "description", "detalle", "nombre", "denominacion")


@dataclass
class HojaLeida:
    nombre: str
    columna_codigo: str
    columna_descripcion: str
    filas: int


@dataclass
class Repuestos:
    descripciones: dict[str, str] = field(default_factory=dict)  # clave(código) -> descripción
    codigos: dict[str, str] = field(default_factory=dict)  # clave(código) -> código tal cual
    hojas: list[HojaLeida] = field(default_factory=list)

    def opciones(self) -> list[str]:
        return sorted(self.codigos.values(), key=clave)

    def codigo(self, codigo: str) -> str:
        """Código tal como figura en el catálogo ('flt-01 ' -> 'FLT-01')."""
        return self.codigos.get(clave(codigo), codigo)

    def descripcion(self, codigo: str) -> str | None:
        return self.descripciones.get(clave(codigo))

    @classmethod
    def desde_bytes(cls, datos: bytes) -> Repuestos:
        repuestos = cls()
        wb = load_workbook(io.BytesIO(datos), data_only=True, read_only=True)
        for ws in wb.worksheets:
            repuestos._leer_hoja(ws.title, [list(f) for f in ws.iter_rows(values_only=True)])
        wb.close()
        return repuestos

    def _leer_hoja(self, nombre: str, filas: list[list[object]]) -> None:
        encabezado = _buscar_encabezado(filas)
        if encabezado is None:
            return
        n_fila, col_cod, col_des = encabezado
        leidas = 0
        for fila in filas[n_fila + 1:]:
            codigo = _texto_codigo(fila[col_cod] if col_cod < len(fila) else None)
            descripcion = limpiar(fila[col_des] if col_des < len(fila) else None)
            if not codigo or not descripcion:
                continue
            leidas += 1
            llave = clave(codigo)
            if llave not in self.descripciones:
                self.descripciones[llave] = descripcion
                self.codigos[llave] = codigo
        self.hojas.append(HojaLeida(
            nombre, limpiar(filas[n_fila][col_cod]), limpiar(filas[n_fila][col_des]), leidas,
        ))


def _columna(fila: list[object], candidatos: tuple[str, ...]) -> int | None:
    """Columna cuyo encabezado coincide con el candidato de mayor prioridad."""
    textos = [clave(limpiar(c)).strip(" :.") for c in fila]
    for candidato in map(clave, candidatos):
        for i, texto in enumerate(textos):
            if texto == candidato or texto.startswith(candidato):
                return i
    return None


def _buscar_encabezado(filas: list[list[object]]) -> tuple[int, int, int] | None:
    """(fila, columna del código, columna de la descripción) o None."""
    for n, fila in enumerate(filas[:_FILAS_PARA_ENCABEZADO]):
        col_cod = _columna(fila, _COLUMNAS_CODIGO)
        col_des = _columna(fila, _COLUMNAS_DESCRIPCION)
        if col_cod is not None and col_des is not None and col_cod != col_des:
            return n, col_cod, col_des
    return None


def _texto_codigo(valor: object) -> str:
    """12345.0 (número en Excel) -> '12345'."""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return limpiar(valor)
